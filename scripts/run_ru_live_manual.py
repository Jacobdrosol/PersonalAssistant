from __future__ import annotations

import argparse
from dataclasses import asdict
from datetime import datetime
import json
from pathlib import Path
import shutil

from openpyxl import load_workbook

from assistant_app.database import Database
from assistant_app.production_log_engine import (
    DateMatchedProductionLogUpdater,
    ProductionLogError,
    SheetImportRule,
    _coerce_for_cell,
    _is_empty,
    _normalize_route,
    _parse_datetime,
    read_csv_rows,
)
from scripts.audit_workbook_preservation import audit


def load_configuration(database_path: Path):
    db = Database(database_path)
    try:
        client = next(item for item in db.get_production_log_clients() if item.name == "RU-PRD")
        automation = next(
            item for item in db.get_production_log_automations(client.id) if item.name == "RU Daily Dry Runs"
        )
        rules = [
            SheetImportRule(
                sheet_name=config.sheet_name,
                data_start_row=config.data_start_row,
                destination_columns=config.column_mappings,
                source_columns=config.source_mappings,
                route_values=config.route_values,
            )
            for config in db.get_production_log_sheet_configs(client.id)
        ]
        return client, automation, rules
    finally:
        db.close()


def build_plan(csv_path: Path, workbook_path: Path, automation, rules: list[SheetImportRule]) -> dict:
    headers, rows = read_csv_rows(csv_path)
    updater = DateMatchedProductionLogUpdater(
        workbook_path,
        routing_column=automation.routing_column,
        source_date_column=automation.source_date_column,
        source_sort_column=automation.source_sort_column,
        target_date_column=automation.target_date_column,
        rules=rules,
    )
    preview = updater.import_rows(headers, rows, apply=False)
    header_lookup = {header.casefold().strip(): header for header in headers}
    routing_header = header_lookup[automation.routing_column.casefold().strip()]
    date_header = header_lookup[automation.source_date_column.casefold().strip()]
    sort_header = header_lookup[automation.source_sort_column.casefold().strip()]
    rule_lookup = {
        _normalize_route(value): rule
        for rule in rules
        for value in (rule.route_values or [rule.sheet_name])
    }
    numbered_rows = [(index + 2, row) for index, row in enumerate(rows)]
    numbered_rows.sort(key=lambda item: (_parse_datetime(item[1][sort_header]) or datetime.max, item[0]))
    writes: list[dict] = []
    issues: list[dict] = []
    pending: set[tuple[str, str]] = set()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=True)
    try:
        for source_row_number, source_row in numbered_rows:
            rule = rule_lookup.get(_normalize_route(source_row.get(routing_header, "")))
            if rule is None:
                issues.append({"source_row_number": source_row_number, "reason": "No route", "source_row": source_row})
                continue
            run_datetime = _parse_datetime(source_row.get(date_header, ""))
            if run_datetime is None:
                issues.append(
                    {"source_row_number": source_row_number, "sheet": rule.sheet_name, "reason": "Invalid date", "source_row": source_row}
                )
                continue
            sheet = workbook[rule.sheet_name]
            target_row = None
            for row_number in range(rule.data_start_row, sheet.max_row + 1):
                value = sheet[f"{automation.target_date_column}{row_number}"].value
                parsed = value if isinstance(value, datetime) else _parse_datetime(value)
                if parsed is not None and parsed.date() == run_datetime.date():
                    target_row = row_number
                    break
            if target_row is None:
                issues.append(
                    {
                        "source_row_number": source_row_number,
                        "sheet": rule.sheet_name,
                        "reason": f"No date row for {run_datetime.date().isoformat()}",
                        "source_row": source_row,
                    }
                )
                continue
            values: dict[str, object] = {}
            current: dict[str, object] = {}
            for source_name, destination in rule.resolved_mappings():
                actual_source = header_lookup[source_name.casefold().strip()]
                cell = sheet[f"{destination}{target_row}"]
                current[destination] = cell.value
                values[destination] = _coerce_for_cell(source_row[actual_source], str(getattr(cell, "number_format", "") or ""))
            coordinates = [(rule.sheet_name, f"{column}{target_row}") for column in values]
            occupied = {column: value for column, value in current.items() if not _is_empty(value)}
            reserved = [coordinate for coordinate in coordinates if coordinate in pending]
            if occupied or reserved:
                issues.append(
                    {
                        "source_row_number": source_row_number,
                        "sheet": rule.sheet_name,
                        "target_row": target_row,
                        "reason": "Destination cells are occupied or already reserved",
                        "current_values": current,
                        "source_row": source_row,
                    }
                )
                continue
            pending.update(coordinates)
            writes.append(
                {
                    "source_row_number": source_row_number,
                    "sheet": rule.sheet_name,
                    "target_row": target_row,
                    "values": values,
                    "source_row": source_row,
                }
            )
    finally:
        workbook.close()
    if preview.rows_written != len(writes) or preview.rows_skipped != len(issues):
        raise ProductionLogError("Independent preview did not agree with the production updater.")
    return {
        "csv": str(csv_path),
        "workbook": str(workbook_path),
        "source_rows": len(rows),
        "writes": writes,
        "issues": issues,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("csv", type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument(
        "--database",
        type=Path,
        default=Path.home() / "AppData/Roaming/PersonalAssistant/assistant.db",
    )
    arguments = parser.parse_args()
    client, automation, rules = load_configuration(arguments.database)
    workbook_path = Path(client.workbook_path or "")
    plan = build_plan(arguments.csv.resolve(), workbook_path, automation, rules)
    if not arguments.apply:
        print(json.dumps(plan, indent=2, default=str))
        return 0

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = Path(__file__).resolve().parents[1] / "test_artifacts" / f"ru_live_run_{stamp}"
    run_dir.mkdir(parents=True, exist_ok=False)
    baseline_path = run_dir / "RU live pre-update.xlsx"
    shutil.copy2(workbook_path, baseline_path)
    updater = DateMatchedProductionLogUpdater(
        workbook_path,
        routing_column=automation.routing_column,
        source_date_column=automation.source_date_column,
        source_sort_column=automation.source_sort_column,
        target_date_column=automation.target_date_column,
        rules=rules,
    )
    result = updater.import_csv(arguments.csv.resolve(), apply=True)
    allowed = {
        (write["sheet"], f"{column}{write['target_row']}")
        for write in plan["writes"]
        for column in write["values"]
    }
    integrity = audit(baseline_path, workbook_path, allowed)
    if not integrity["formatting_preserved"]:
        shutil.copy2(baseline_path, workbook_path)
        integrity["restored_from_baseline"] = True
        (run_dir / "FAILED_integrity_report.json").write_text(
            json.dumps(integrity, indent=2, default=str), encoding="utf-8"
        )
        raise ProductionLogError("Integrity audit failed; the pre-update workbook was restored.")
    summary = {
        "plan": plan,
        "result": {
            "source_rows": result.source_rows,
            "rows_written": result.rows_written,
            "cells_written": result.cells_written,
            "rows_skipped": result.rows_skipped,
            "issues": [asdict(issue) for issue in result.issues],
        },
        "integrity": integrity,
        "pre_update_copy": str(baseline_path),
    }
    report_path = run_dir / "live_update_report.json"
    report_path.write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    print(json.dumps({"run_dir": str(run_dir), **summary["result"], "formatting_preserved": True}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
