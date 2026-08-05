from __future__ import annotations

import argparse
from hashlib import sha256
import json
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook


FORMAT_ARCHIVE_PREFIXES = (
    "xl/styles.xml",
    "xl/theme/",
    "xl/drawings/",
    "xl/charts/",
    "xl/media/",
    "xl/tables/",
    "xl/comments",
    "xl/threadedComments/",
    "xl/persons/",
)


def _hashes(path: Path) -> dict[str, str]:
    with ZipFile(path) as archive:
        return {
            name: sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name == "xl/styles.xml" or name.startswith(FORMAT_ARCHIVE_PREFIXES[1:])
        }


def _cell_style(cell) -> tuple[Any, ...]:
    return (
        bool(cell.has_style),
        tuple(cell._style) if cell.has_style else None,
        cell.number_format,
        cell.quotePrefix,
        cell.pivotButton,
    )


def _row_dimension(dimension) -> tuple[Any, ...]:
    return (
        dimension.height,
        dimension.hidden,
        dimension.outlineLevel,
        dimension.collapsed,
        dimension.thickTop,
        dimension.thickBot,
        dimension.style_id,
        dimension.customFormat,
        dimension.customHeight,
    )


def _column_dimension(dimension) -> tuple[Any, ...]:
    return (
        dimension.min,
        dimension.max,
        dimension.width,
        dimension.hidden,
        dimension.bestFit,
        dimension.outlineLevel,
        dimension.collapsed,
        dimension.style_id,
        dimension.customWidth,
    )


def _xml_bytes(node) -> bytes:
    return ET.tostring(node.to_tree(), encoding="utf-8")


def _conditional_formatting(sheet) -> list[tuple[str, list[str]]]:
    result: list[tuple[str, list[str]]] = []
    for conditional_format, rules in sheet.conditional_formatting._cf_rules.items():
        result.append((str(conditional_format.sqref), [_xml_bytes(rule).decode("utf-8") for rule in rules]))
    return result


def _data_validations(sheet) -> list[str]:
    collection = sheet.data_validations
    if collection is None:
        return []
    return [_xml_bytes(item).decode("utf-8") for item in collection.dataValidation]


def audit(baseline_path: Path, updated_path: Path, allowed_cells: set[tuple[str, str]]) -> dict[str, Any]:
    differences: list[dict[str, Any]] = []

    baseline_archive = _hashes(baseline_path)
    updated_archive = _hashes(updated_path)
    for name in sorted(set(baseline_archive) | set(updated_archive)):
        if baseline_archive.get(name) != updated_archive.get(name):
            differences.append({"type": "format_archive_part", "part": name})

    baseline = load_workbook(baseline_path, data_only=False, keep_links=True)
    updated = load_workbook(updated_path, data_only=False, keep_links=True)
    try:
        if baseline.sheetnames != updated.sheetnames:
            differences.append(
                {"type": "sheet_order", "before": baseline.sheetnames, "after": updated.sheetnames}
            )
        for sheet_name in baseline.sheetnames:
            if sheet_name not in updated.sheetnames:
                continue
            before = baseline[sheet_name]
            after = updated[sheet_name]

            before_cells = set(before._cells)
            after_cells = set(after._cells)
            unexpected_cells = {
                coordinate
                for coordinate in before_cells ^ after_cells
                if (sheet_name, before.cell(*coordinate).coordinate) not in allowed_cells
            }
            if unexpected_cells:
                differences.append(
                    {"type": "cell_structure", "sheet": sheet_name, "cells": sorted(unexpected_cells)}
                )

            for row, column in sorted(before_cells | after_cells):
                before_cell = before.cell(row, column)
                after_cell = after.cell(row, column)
                coordinate = before_cell.coordinate
                if _cell_style(before_cell) != _cell_style(after_cell):
                    differences.append(
                        {"type": "cell_style", "sheet": sheet_name, "cell": coordinate}
                    )
                if (sheet_name, coordinate) not in allowed_cells:
                    if before_cell.value != after_cell.value:
                        differences.append(
                            {
                                "type": "cell_value",
                                "sheet": sheet_name,
                                "cell": coordinate,
                                "before": before_cell.value,
                                "after": after_cell.value,
                            }
                        )
                    if before_cell.comment != after_cell.comment:
                        differences.append({"type": "cell_comment", "sheet": sheet_name, "cell": coordinate})
                    before_link = before_cell.hyperlink.target if before_cell.hyperlink else None
                    after_link = after_cell.hyperlink.target if after_cell.hyperlink else None
                    if before_link != after_link:
                        differences.append({"type": "cell_hyperlink", "sheet": sheet_name, "cell": coordinate})

            before_rows = {key: _row_dimension(value) for key, value in before.row_dimensions.items()}
            after_rows = {key: _row_dimension(value) for key, value in after.row_dimensions.items()}
            if before_rows != after_rows:
                differences.append({"type": "row_dimensions", "sheet": sheet_name})
            before_columns = {key: _column_dimension(value) for key, value in before.column_dimensions.items()}
            after_columns = {key: _column_dimension(value) for key, value in after.column_dimensions.items()}
            if before_columns != after_columns:
                differences.append({"type": "column_dimensions", "sheet": sheet_name})
            if sorted(str(item) for item in before.merged_cells.ranges) != sorted(
                str(item) for item in after.merged_cells.ranges
            ):
                differences.append({"type": "merged_cells", "sheet": sheet_name})
            if _conditional_formatting(before) != _conditional_formatting(after):
                differences.append({"type": "conditional_formatting", "sheet": sheet_name})
            if _data_validations(before) != _data_validations(after):
                differences.append({"type": "data_validations", "sheet": sheet_name})
            for property_name in (
                "freeze_panes",
                "sheet_format",
                "sheet_properties",
                "page_margins",
                "page_setup",
                "print_options",
            ):
                if str(getattr(before, property_name)) != str(getattr(after, property_name)):
                    differences.append(
                        {"type": property_name, "sheet": sheet_name}
                    )
            if before.auto_filter.ref != after.auto_filter.ref:
                differences.append({"type": "auto_filter", "sheet": sheet_name})
            if sorted(before.tables) != sorted(after.tables):
                differences.append({"type": "tables", "sheet": sheet_name})
            if len(before._charts) != len(after._charts):
                differences.append({"type": "chart_count", "sheet": sheet_name})
            if len(before._images) != len(after._images):
                differences.append({"type": "image_count", "sheet": sheet_name})
    finally:
        baseline.close()
        updated.close()

    return {
        "baseline": str(baseline_path),
        "updated": str(updated_path),
        "allowed_value_changes": [f"{sheet}!{cell}" for sheet, cell in sorted(allowed_cells)],
        "formatting_preserved": not differences,
        "difference_count": len(differences),
        "differences": differences,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("baseline", type=Path)
    parser.add_argument("updated", type=Path)
    parser.add_argument("--allow", action="append", default=[])
    parser.add_argument("--report", type=Path, required=True)
    arguments = parser.parse_args()
    allowed: set[tuple[str, str]] = set()
    for item in arguments.allow:
        sheet, separator, coordinate = item.rpartition("!")
        if not separator:
            raise SystemExit(f"Invalid allowed cell: {item}")
        allowed.add((sheet, coordinate.upper()))
    report = audit(arguments.baseline, arguments.updated, allowed)
    arguments.report.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    print(json.dumps(report, indent=2, default=str))
    return 0 if report["formatting_preserved"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
