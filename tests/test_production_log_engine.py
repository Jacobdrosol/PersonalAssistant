from __future__ import annotations

from pathlib import Path
from datetime import datetime
import tempfile
import unittest
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from assistant_app.production_log_engine import (
    DateMatchedProductionLogUpdater,
    OutlookCsvSource,
    ProductionLogUpdater,
    SheetImportRule,
    read_csv_bytes,
)


class ClassicOutlookSynchronizationTests(unittest.TestCase):
    def test_starts_every_configured_send_receive_group(self) -> None:
        class Group:
            def __init__(self) -> None:
                self.started = False

            def Start(self) -> None:
                self.started = True

        groups = [Group(), Group()]

        class SyncObjects:
            Count = 2

            @staticmethod
            def Item(index: int):
                return groups[index - 1]

        class Namespace:
            Offline = False

        namespace = Namespace()
        namespace.SyncObjects = SyncObjects()

        OutlookCsvSource._start_outlook_sync(namespace)

        self.assertTrue(all(group.started for group in groups))


class ProductionLogUpdaterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "production.xlsx"
        workbook = Workbook()
        mailed = workbook.active
        mailed.title = "Mailed"
        emailed = workbook.create_sheet("Emailed")
        for sheet in (mailed, emailed):
            sheet["A1"] = "Name"
            sheet["B1"] = "Count"
            sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="00FF00")
            sheet["B2"].fill = PatternFill(fill_type="solid", fgColor="00FF00")
        mailed["A2"] = "existing"
        workbook.save(self.workbook_path)
        workbook.close()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _updater(self) -> ProductionLogUpdater:
        rules = [
            SheetImportRule(
                sheet_name="Mailed",
                data_start_row=2,
                destination_columns={"name": "A", "count": "B"},
                source_columns={"name": "Customer", "count": "Total"},
                route_values=["MAIL"],
            ),
            SheetImportRule(
                sheet_name="Emailed",
                data_start_row=2,
                destination_columns={"name": "A", "count": "B"},
                source_columns={"name": "Customer", "count": "Total"},
                route_values=["EMAIL"],
            ),
        ]
        return ProductionLogUpdater(self.workbook_path, "Channel", rules)

    def test_routes_rows_preserves_existing_values_and_copies_format(self) -> None:
        rows = [
            {"Channel": "MAIL", "Customer": "Alpha", "Total": "12"},
            {"Channel": "EMAIL", "Customer": "Beta", "Total": "7"},
            {"Channel": "UNKNOWN", "Customer": "Gamma", "Total": "3"},
        ]
        result = self._updater().import_rows(["Channel", "Customer", "Total"], rows)

        self.assertEqual(result.rows_written, 2)
        self.assertEqual(result.cells_written, 4)
        self.assertEqual(result.unrouted_rows, 1)
        workbook = load_workbook(self.workbook_path)
        try:
            mailed = workbook["Mailed"]
            emailed = workbook["Emailed"]
            self.assertEqual(mailed["A2"].value, "existing")
            self.assertIsNone(mailed["B2"].value)
            self.assertEqual(mailed["A3"].value, "Alpha")
            self.assertEqual(mailed["B3"].value, 12)
            self.assertEqual(mailed["A3"].fill.fgColor.rgb, mailed["A2"].fill.fgColor.rgb)
            self.assertEqual(emailed["A2"].value, "Beta")
            self.assertEqual(emailed["B2"].value, 7)
        finally:
            workbook.close()

    def test_second_import_appends_without_overwriting_first_import(self) -> None:
        updater = self._updater()
        updater.import_rows(
            ["Channel", "Customer", "Total"],
            [{"Channel": "EMAIL", "Customer": "First", "Total": "1"}],
        )
        updater.import_rows(
            ["Channel", "Customer", "Total"],
            [{"Channel": "EMAIL", "Customer": "Second", "Total": "2"}],
        )
        workbook = load_workbook(self.workbook_path, data_only=True)
        try:
            sheet = workbook["Emailed"]
            self.assertEqual(sheet["A2"].value, "First")
            self.assertEqual(sheet["B2"].value, 1)
            self.assertEqual(sheet["A3"].value, "Second")
            self.assertEqual(sheet["B3"].value, 2)
        finally:
            workbook.close()

    def test_reads_utf8_bom_csv_with_header_row(self) -> None:
        headers, rows = read_csv_bytes("Channel,Customer,Total\r\nMAIL,Alpha,12\r\n".encode("utf-8-sig"))
        self.assertEqual(headers, ["Channel", "Customer", "Total"])
        self.assertEqual(rows[0]["Customer"], "Alpha")


class DateMatchedProductionLogUpdaterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "dated-production.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Auto - Apply"
        sheet.append(["Run Date", "Dry Run ID", "Dry Run Counts", "Notes"])
        sheet.append([datetime(2026, 8, 4), None, None, "available"])
        sheet.append([datetime(2026, 8, 5), 999, None, "already populated"])
        sheet["B2"].fill = PatternFill(fill_type="solid", fgColor="00FFFF00")
        sheet["C2"] = ""
        sheet["C2"].fill = PatternFill(fill_type="solid", fgColor="00FFFF00")
        workbook.save(self.workbook_path)
        workbook.close()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_matches_date_requires_both_destinations_empty_and_does_not_fall_through(self) -> None:
        rule = SheetImportRule(
            sheet_name="Auto - Apply",
            data_start_row=2,
            destination_columns={"id": "B", "count": "C"},
            source_columns={"id": "Distribution Run ID", "count": "Record Count"},
            route_values=["NUB_Single_Auto_Renewal_Dry_Apply"],
        )
        updater = DateMatchedProductionLogUpdater(
            self.workbook_path,
            routing_column="Select Set Name",
            source_date_column="Beginning Date/Time",
            source_sort_column="Beginning Date/Time",
            target_date_column="A",
            rules=[rule],
            backup_root=Path(self.temp_dir.name) / "backups",
        )
        with ZipFile(self.workbook_path) as archive:
            worksheet_xml = archive.read("xl/worksheets/sheet1.xml")
            original_parts = {name: archive.read(name) for name in archive.namelist()}
        self.assertIn(b'r="C2"', worksheet_xml, worksheet_xml.decode("utf-8"))
        headers = [
            "Distribution Run ID",
            "Select Set Name",
            "Beginning Date/Time",
            "Record Count",
        ]
        rows = [
            {
                "Distribution Run ID": "3000",
                "Select Set Name": "NUB_Single_Auto_Renewal_Dry_Apply",
                "Beginning Date/Time": "8/6/2026 4:00:00 AM",
                "Record Count": "4",
            },
            {
                "Distribution Run ID": "2685",
                "Select Set Name": "NUB_Single_Auto_Renewal_Dry_Apply",
                "Beginning Date/Time": "8/4/2026 4:47:02 AM",
                "Record Count": "3",
            },
            {
                "Distribution Run ID": "2999",
                "Select Set Name": "NUB_Single_Auto_Renewal_Dry_Apply",
                "Beginning Date/Time": "8/5/2026 4:00:00 AM",
                "Record Count": "5",
            },
        ]

        result = updater.import_rows(headers, rows)

        self.assertEqual(result.rows_written, 1)
        self.assertEqual(result.cells_written, 2)
        self.assertEqual(result.rows_skipped, 2)
        self.assertEqual([issue.source_row_number for issue in result.issues], [4, 2])
        self.assertEqual(result.issues[0].target_row_number, 3)
        self.assertEqual(result.issues[0].target_values, {"B": 999, "C": None})
        self.assertIsNone(result.issues[1].target_row_number)
        workbook = load_workbook(self.workbook_path, data_only=True)
        try:
            sheet = workbook["Auto - Apply"]
            self.assertEqual(sheet["B2"].value, 2685)
            self.assertEqual(sheet["C2"].value, 3)
            self.assertEqual(sheet["B3"].value, 999)
            self.assertIsNone(sheet["C3"].value)
            self.assertEqual(sheet["B2"].fill.fill_type, "solid")
            self.assertTrue(str(sheet["B2"].fill.fgColor.rgb).endswith("FFFF00"))
        finally:
            workbook.close()
        with ZipFile(self.workbook_path) as archive:
            self.assertEqual(list(original_parts), archive.namelist())
            for name, original_data in original_parts.items():
                if name not in {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}:
                    self.assertEqual(original_data, archive.read(name), name)
            workbook_xml = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            calculation = workbook_xml.find(f"{{{namespace}}}calcPr")
            self.assertIsNotNone(calculation)
            self.assertEqual(calculation.attrib.get("calcMode"), "auto")
            self.assertEqual(calculation.attrib.get("fullCalcOnLoad"), "1")
            self.assertEqual(calculation.attrib.get("forceFullCalc"), "1")


if __name__ == "__main__":
    unittest.main()
